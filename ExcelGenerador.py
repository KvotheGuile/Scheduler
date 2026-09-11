import json
import re
from datetime import datetime, time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
THIN_BORDER = Border(*(Side(style="thin", color="B7B7B7"),) * 4)
FILL_MANUAL = PatternFill("solid", fgColor="FFF2CC")
SUBTITULO_FONT = Font(name=FONT_NAME, bold=True, italic=True, size=11)
NOTA_FONT = Font(name=FONT_NAME, italic=True, size=9, color="555555")

DIAS_ORDEN = ["Lu", "Ma", "Mi", "Ju", "Vi", "Sa", "Do"]
DIA_REGEX = re.compile("|".join(DIAS_ORDEN))

HORA_INICIO_GRID = time(7, 0)
HORA_FIN_GRID = time(22, 0)
PASO_MIN = 30

NUM_CLUSTERS = 3
CLUSTER_TITULOS = ["Semanas 1-5", "Semanas 6-10", "Semanas 11-15"]
PERIODO_A_CLUSTERS = {
    "1": (0,), "2": (1,), "3": (2,),
    "4": (0, 1), "5": (1, 2), "6": (0, 1, 2),
}

PALETA = [
    "B4C6E7", "F4B183", "C6E0B4", "FFD966", "D9B3FF",
    "9DC3E6", "F8CBAD", "A9D18E", "FFE699", "B4A7D6",
    "8FAADC", "F4CCCC", "D9EAD3", "FCE5CD", "C9DAF8",
]

PRINCIPAL_LAYOUT = [
    ("periodo_escolar", "Semestre"),
    ("modelo", "Modelo"),
    ("programa", "Programa"),
    ("semestre", "Semestre"),
    ("clave", "Clave"),
    ("modulo", "Unidad de formación/Módulo"),
    ("udca", "UdCA"),
    ("clasificacion", "Clasificación"),
    ("pronostico", "Pronóstico"),
    ("grupo", "Grupo"),
    ("modalidad", "Modalidad"),
    ("idioma", "Idioma"),
    ("departamento", "Departamento"),
    ("roster", "% Roster"),
    ("nomina", "Nómina"),
    ("profesor", "Profesor"),
    ("udceq", "UdCeq"),
    ("periodo", "Período"),
    ("dia", "Día"),
    ("hora_inicio", "Hora inicio"),
    ("hora_fin", "Hora fin"),
    ("salon", "Salón"),
    ("horas", "Horas (aux.)"),
    ("clave_horario", "Clave horario (aux.)"),
]
COL = {clave: i + 1 for i, (clave, _) in enumerate(PRINCIPAL_LAYOUT)}

TABLA_SECCION_COLUMNAS = [
    ("Período", None),
    ("Grupo", None),
    ("Clave", None),
    ("Módulo", "modulo"),
    ("Profesor", "profesor"),
    ("Salón", "salon"),
    ("Días", "dia"),
    ("Hora inicio", "hora_inicio"),
    ("Hora fin", "hora_fin"),
    ("Carga (UdCeq)", "udceq"),
]


def slot_index(hora):
    minutos = (hora.hour * 60 + hora.minute) - (HORA_INICIO_GRID.hour * 60 + HORA_INICIO_GRID.minute)
    return minutos // PASO_MIN


def formula_lookup(clave_col, clave_valor, rango_clave, ultima_fila):
    col_letra = get_column_letter(COL[clave_col])
    return (
        f'=IFERROR(INDEX(Principal!${col_letra}$2:${col_letra}${ultima_fila},'
        f'MATCH("{clave_valor}",{rango_clave},0)),"")'
    )


def cargar_datos(ruta_json):
    with open(ruta_json, encoding="utf-8") as f:
        datos = json.load(f)
    for d in datos:
        d["grupo"] = str(d["grupo"])
        d["horaInicio"] = datetime.strptime(d["horaInicio"].strip(), "%H:%M").time()
        d["horaFinal"] = datetime.strptime(d["horaFinal"].strip(), "%H:%M").time()
        d["carga"] = float(d.get("carga", 0))
    return datos


def escribir_principal(wb, datos):
    ws = wb.create_sheet("Principal")

    for c, (_clave, titulo) in enumerate(PRINCIPAL_LAYOUT, start=1):
        celda = ws.cell(row=1, column=c, value=titulo)
        celda.font = HEADER_FONT
        celda.fill = HEADER_FILL
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
        celda.border = THIN_BORDER

    for r, clase in enumerate(datos, start=2):
        clasificacion = "Módulo" if re.search(r"B\.\d+$", clase["claseId"]) else "Materia"
        valores = {
            "modelo": "Tec21",
            "programa": clase["carrera"],
            "semestre": clase["semestre"],
            "clave": clase["claseId"],
            "modulo": clase["claseNombre"],
            "clasificacion": clasificacion,
            "grupo": clase["grupo"],
            "modalidad": "CPRS",
            "idioma": "Español",
            "nomina": clase["profeId"],
            "profesor": clase["profeNombre"],
            "udceq": clase["carga"],
            "periodo": int(clase["grupo"][0]),
            "dia": clase["dias"],
            "hora_inicio": clase["horaInicio"],
            "hora_fin": clase["horaFinal"],
            "salon": clase["salon"],
        }
        for clave_col, valor in valores.items():
            celda = ws.cell(row=r, column=COL[clave_col], value=valor)
            celda.font = Font(name=FONT_NAME)
            celda.border = THIN_BORDER
            if clave_col in ("hora_inicio", "hora_fin"):
                celda.number_format = "hh:mm"

        ws.cell(row=r, column=COL["periodo_escolar"]).fill = FILL_MANUAL
        ws.cell(row=r, column=COL["periodo_escolar"]).border = THIN_BORDER

        f_ini = f"{get_column_letter(COL['hora_inicio'])}{r}"
        f_fin = f"{get_column_letter(COL['hora_fin'])}{r}"
        celda = ws.cell(row=r, column=COL["horas"], value=f"=({f_fin}-{f_ini})*24")
        celda.font = Font(name=FONT_NAME)
        celda.border = THIN_BORDER

        g = f"{get_column_letter(COL['grupo'])}{r}"
        cid = f"{get_column_letter(COL['clave'])}{r}"
        celda = ws.cell(row=r, column=COL["clave_horario"], value=f'={g}&"-"&{cid}')
        celda.font = Font(name=FONT_NAME)
        celda.border = THIN_BORDER

    anchos = {
        "periodo_escolar": 10, "modelo": 9, "programa": 10, "semestre": 9,
        "clave": 12, "modulo": 26, "udca": 7, "clasificacion": 13,
        "pronostico": 11, "grupo": 8, "modalidad": 10, "idioma": 9,
        "departamento": 13, "roster": 9, "nomina": 10, "profesor": 16,
        "udceq": 8, "periodo": 9, "dia": 10, "hora_inicio": 11,
        "hora_fin": 10, "salon": 10, "horas": 10, "clave_horario": 16,
    }
    for clave_col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(COL[clave_col])].width = ancho

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(PRINCIPAL_LAYOUT))}{len(datos) + 1}"
    return ws


def escribir_resumen_cargas(wb, datos, ultima_fila):
    ws = wb.create_sheet("Resumen de cargas")

    headers = ["Nómina", "Profesor", "No. de grupos", "Carga total (UdCeq)", "Horas totales"]
    for c, titulo in enumerate(headers, start=1):
        celda = ws.cell(row=1, column=c, value=titulo)
        celda.font = HEADER_FONT
        celda.fill = HEADER_FILL
        celda.border = THIN_BORDER

    profesores = {}
    for d in datos:
        profesores.setdefault(d["profeId"], d["profeNombre"])

    col_nomina = get_column_letter(COL["nomina"])
    col_udceq = get_column_letter(COL["udceq"])
    col_horas = get_column_letter(COL["horas"])
    rango_id = f"Principal!${col_nomina}$2:${col_nomina}${ultima_fila}"
    rango_carga = f"Principal!${col_udceq}$2:${col_udceq}${ultima_fila}"
    rango_horas = f"Principal!${col_horas}$2:${col_horas}${ultima_fila}"

    for r, (profe_id, nombre) in enumerate(profesores.items(), start=2):
        ws.cell(row=r, column=1, value=profe_id).font = Font(name=FONT_NAME)
        ws.cell(row=r, column=2, value=nombre).font = Font(name=FONT_NAME)
        clave = f"A{r}"
        ws.cell(row=r, column=3, value=f"=COUNTIF({rango_id},{clave})").font = Font(name=FONT_NAME)
        ws.cell(row=r, column=4, value=f"=SUMIF({rango_id},{clave},{rango_carga})").font = Font(name=FONT_NAME)
        ws.cell(row=r, column=5, value=f"=SUMIF({rango_id},{clave},{rango_horas})").font = Font(name=FONT_NAME)
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = THIN_BORDER

    for col, ancho in zip("ABCDE", (10, 18, 14, 18, 14)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    return ws


def escribir_hoja_semestre(wb, carrera, semestre, clases, ultima_fila):
    ws = wb.create_sheet(f"{carrera} Sem {semestre}"[:31])

    col_clave_h = get_column_letter(COL["clave_horario"])
    rango_clave = f"Principal!${col_clave_h}$2:${col_clave_h}${ultima_fila}"

    ws.cell(row=1, column=1, value="Carrera:").font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=1, column=2, value=carrera).font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=1, column=4, value="Semestre:").font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=1, column=5, value=semestre).font = Font(name=FONT_NAME, bold=True)

    secciones = {}
    for clase in clases:
        periodo, seccion = clase["grupo"][0], clase["grupo"][1:]
        secciones.setdefault(seccion, {}).setdefault(periodo, []).append(clase)

    dias_presentes = DIAS_ORDEN[:5]
    ancho_bloque = len(dias_presentes)
    espacio = 1
    total_slots = slot_index(HORA_FIN_GRID)
    col_modulo = get_column_letter(COL["modulo"])
    col_profe = get_column_letter(COL["profesor"])
    col_salon = get_column_letter(COL["salon"])

    fila = 3
    for seccion in sorted(secciones):
        periodos_seccion = secciones[seccion]
        periodos_ordenados = sorted(periodos_seccion)

        ws.cell(row=fila, column=1, value=f"Sección {seccion}").font = SUBTITULO_FONT
        resumen = "   |   ".join(
            f"Período {p}: Grupo {periodos_seccion[p][0]['grupo']}" for p in periodos_ordenados
        )
        ws.cell(row=fila, column=2, value=resumen).font = NOTA_FONT
        fila += 1

        fila_tabla_header = fila
        for c, (titulo, _clave) in enumerate(TABLA_SECCION_COLUMNAS, start=1):
            celda = ws.cell(row=fila_tabla_header, column=c, value=titulo)
            celda.font = HEADER_FONT
            celda.fill = HEADER_FILL
            celda.alignment = Alignment(horizontal="center", wrap_text=True)
            celda.border = THIN_BORDER

        fila_dato = fila_tabla_header + 1
        for periodo in periodos_ordenados:
            for clase in periodos_seccion[periodo]:
                clave_valor = f"{clase['grupo']}-{clase['claseId']}"
                for c, (_titulo, clave_col) in enumerate(TABLA_SECCION_COLUMNAS, start=1):
                    if clave_col is None:
                        valor = {1: periodo, 2: clase["grupo"], 3: clase["claseId"]}[c]
                        celda = ws.cell(row=fila_dato, column=c, value=valor)
                    else:
                        celda = ws.cell(
                            row=fila_dato, column=c,
                            value=formula_lookup(clave_col, clave_valor, rango_clave, ultima_fila),
                        )
                    celda.font = Font(name=FONT_NAME, size=10)
                    celda.border = THIN_BORDER
                    if clave_col in ("hora_inicio", "hora_fin"):
                        celda.number_format = "hh:mm"
                fila_dato += 1

        fila = fila_dato + 1

        fila_grid_titulo = fila
        fila_grid_header = fila_grid_titulo + 1
        fila_inicio_grid = fila_grid_header + 1

        ws.cell(row=fila_grid_header, column=1, value="Hora").font = HEADER_FONT
        ws.cell(row=fila_grid_header, column=1).fill = HEADER_FILL
        for s in range(total_slots):
            minutos = s * PASO_MIN
            h = HORA_INICIO_GRID.hour + minutos // 60
            m = (HORA_INICIO_GRID.minute + minutos) % 60
            celda = ws.cell(row=fila_inicio_grid + s, column=1, value=time(h, m))
            celda.number_format = "hh:mm"
            celda.font = Font(name=FONT_NAME, size=9)
            celda.border = THIN_BORDER
            celda.fill = PatternFill("solid", fgColor="F2F2F2")

        col_inicio_por_cluster = []
        for cl in range(NUM_CLUSTERS):
            col_inicio = 2 + cl * (ancho_bloque + espacio)
            col_inicio_por_cluster.append(col_inicio)

            celda_titulo = ws.cell(row=fila_grid_titulo, column=col_inicio, value=CLUSTER_TITULOS[cl])
            celda_titulo.font = HEADER_FONT
            celda_titulo.fill = HEADER_FILL
            celda_titulo.alignment = Alignment(horizontal="center")
            ws.merge_cells(
                start_row=fila_grid_titulo, start_column=col_inicio,
                end_row=fila_grid_titulo, end_column=col_inicio + ancho_bloque - 1,
            )

            for j, dia in enumerate(dias_presentes):
                celda = ws.cell(row=fila_grid_header, column=col_inicio + j, value=dia)
                celda.font = HEADER_FONT
                celda.fill = HEADER_FILL
                celda.alignment = Alignment(horizontal="center")

            for s in range(total_slots):
                for j in range(ancho_bloque):
                    ws.cell(row=fila_inicio_grid + s, column=col_inicio + j).border = THIN_BORDER

        for periodo in periodos_ordenados:
            clusters_periodo = PERIODO_A_CLUSTERS.get(str(periodo), tuple(range(NUM_CLUSTERS)))
            for clase in periodos_seccion[periodo]:
                dias = [d for d in DIAS_ORDEN if d in DIA_REGEX.findall(clase["dias"])]
                s_ini = max(0, slot_index(clase["horaInicio"]))
                s_fin = min(total_slots, slot_index(clase["horaFinal"]))
                if s_fin <= s_ini:
                    continue

                color = PALETA[sum(ord(c) for c in clase["claseId"]) % len(PALETA)]
                clave_valor = f"{clase['grupo']}-{clase['claseId']}"
                formula = (
                    f'=IFERROR(INDEX(Principal!${col_modulo}$2:${col_modulo}${ultima_fila},'
                    f'MATCH("{clave_valor}",{rango_clave},0))&CHAR(10)&'
                    f'INDEX(Principal!${col_profe}$2:${col_profe}${ultima_fila},'
                    f'MATCH("{clave_valor}",{rango_clave},0))&CHAR(10)&'
                    f'INDEX(Principal!${col_salon}$2:${col_salon}${ultima_fila},'
                    f'MATCH("{clave_valor}",{rango_clave},0)),"")'
                )

                fila_top = fila_inicio_grid + s_ini
                fila_bottom = fila_inicio_grid + s_fin - 1

                for cl in clusters_periodo:
                    col_inicio = col_inicio_por_cluster[cl]
                    for dia in dias:
                        if dia not in dias_presentes:
                            continue
                        col = col_inicio + dias_presentes.index(dia)

                        celda = ws.cell(row=fila_top, column=col, value=formula)
                        celda.fill = PatternFill("solid", fgColor=color)
                        celda.font = Font(name=FONT_NAME, size=9)
                        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        if fila_bottom > fila_top:
                            ws.merge_cells(start_row=fila_top, start_column=col, end_row=fila_bottom, end_column=col)
                        for f in range(fila_top, fila_bottom + 1):
                            ws.cell(row=f, column=col).border = THIN_BORDER
                            ws.cell(row=f, column=col).fill = PatternFill("solid", fgColor=color)

        for s in range(total_slots):
            ws.row_dimensions[fila_inicio_grid + s].height = 15

        fila = fila_inicio_grid + total_slots + 2

    ws.column_dimensions["A"].width = 9
    for i in range(NUM_CLUSTERS):
        col_inicio = 2 + i * (ancho_bloque + espacio)
        for j in range(ancho_bloque):
            ws.column_dimensions[get_column_letter(col_inicio + j)].width = 20
        if i < NUM_CLUSTERS - 1:
            ws.column_dimensions[get_column_letter(col_inicio + ancho_bloque)].width = 3

    ws.freeze_panes = "B3"
    return ws


def generar(ruta_json, ruta_salida):
    datos = cargar_datos(ruta_json)
    if not datos:
        raise SystemExit("El JSON no tiene elementos.")

    wb = Workbook()
    wb.remove(wb.active)

    escribir_principal(wb, datos)
    ultima_fila = len(datos) + 1
    escribir_resumen_cargas(wb, datos, ultima_fila)

    semestres = {}
    for d in datos:
        semestres.setdefault((d["carrera"], d["semestre"]), []).append(d)

    for (carrera, semestre), clases in sorted(semestres.items(), key=lambda kv: kv[0]):
        escribir_hoja_semestre(wb, carrera, semestre, clases, ultima_fila)

    wb.save(ruta_salida)
    print(f"Listo: {ruta_salida} ({len(datos)} clases, {len(semestres)} hoja(s))")

if __name__ == "main":
    nombre_json = "schedule.json"
    nombre_salida = "salida.xlsx"
    generar(nombre_json, nombre_salida)