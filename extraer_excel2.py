
#Para usar: python extraer_excel.py archivo.xlsx --salida data

import argparse
import json
import math
import re
import unicodedata
from datetime import datetime, time
from pathlib import Path

from openpyxl import load_workbook


DIAS = ("Lu", "Ma", "Mi", "Ju", "Vi", "Sa", "Do")


def normalizar(valor):
    texto = unicodedata.normalize("NFKD", str(valor).strip().lower())
    return "".join(c for c in texto if c.isalnum() and not unicodedata.combining(c))


def filas(libro, nombre, columnas):
    #Lee solo columnas declaradas.
    hojas = {normalizar(s.title): s for s in libro}
    if normalizar(nombre) not in hojas:
        raise ValueError(f"Falta la hoja {nombre!r}")
    hoja = hojas[normalizar(nombre)]
    encabezados = {normalizar(c.value): c.column for c in hoja[1] if c.value}
    faltantes = [c for c in columnas if normalizar(c) not in encabezados]
    if faltantes:
        raise ValueError(f"{nombre}: faltan columnas {faltantes}")
    indices = [encabezados[normalizar(c)] for c in columnas]
    for numero, fila in enumerate(hoja.iter_rows(min_row=2), 2):
        celdas = [fila[i - 1] for i in indices]
        if all(c.value is None or c.value == "" for c in celdas):
            continue
        datos = {}
        for campo, celda in zip(columnas, celdas):
            valor = celda.value.strip() if isinstance(celda.value, str) else celda.value
            if valor is None or valor == "" or celda.data_type == "e":
                raise ValueError(
                    f"{nombre}!{celda.coordinate}: {campo} vacío o con error. "
                    "Si contiene una fórmula, recalcula y guarda el Excel."
                )
            datos[campo] = valor
        yield f"{nombre}, fila {numero}", datos


def numero(valor, contexto, entero=False, minimo=0):
    try:
        resultado = float(valor)
    except (ValueError, TypeError):
        raise ValueError(f"{contexto}: número inválido {valor!r}") from None
    if (isinstance(valor, bool) or not math.isfinite(resultado)
            or resultado < minimo or (entero and not resultado.is_integer())):
        raise ValueError(f"{contexto}: número fuera de rango {valor!r}")
    return int(resultado) if entero else resultado


def periodos(valor, contexto):
    #Conserva códigos 1..6; admite uno o varios separados por coma.
    partes = str(valor).split(",")
    resultado = sorted({numero(p.strip(), contexto, entero=True, minimo=1) for p in partes})
    if any(p > 6 for p in resultado):
        raise ValueError(f"{contexto}: los períodos deben estar entre 1 y 6")
    return resultado


def dias(valor):
    texto = re.sub(r"[\s,;]+", "", str(valor)).lower()
    partes = re.findall(r"lu|ma|mi|ju|vi|sa|do", texto)
    if not partes or "".join(partes) != texto:
        raise ValueError(f"Días inválidos: {valor!r}; usa LuMaMiJuViSaDo")
    return [d for d in DIAS if d.lower() in partes]


def minutos(valor):
    if isinstance(valor, datetime):
        valor = valor.time()
    if isinstance(valor, time):
        if valor.second or valor.microsecond:
            raise ValueError(f"La hora debe tener precisión de minutos: {valor}")
        return valor.hour * 60 + valor.minute
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        if not 0 <= valor < 1:
            raise ValueError(f"Fracción de día Excel inválida: {valor}")
        return round(valor * 1440)
    coincidencia = re.fullmatch(r"(\d{1,2}):(\d{2})", str(valor).strip())
    if coincidencia:
        h, m = map(int, coincidencia.groups())
        if 0 <= h < 24 and 0 <= m < 60:
            return h * 60 + m
    raise ValueError(f"Hora inválida: {valor!r}; usa HH:MM")


def restar(intervalos, inicio, fin):
    resultado = []
    for a, b in intervalos:
        if fin <= a or inicio >= b:
            resultado.append((a, b))
        else:
            if a < inicio:
                resultado.append((a, inicio))
            if fin < b:
                resultado.append((fin, b))
    return resultado


def hora(minuto):
    return f"{minuto // 60:02d}:{minuto % 60:02d}"


def extraer(libro, dias_semana=DIAS):
    materias, carreras = [], []
    solicitudes = set()
    columnas = ("Programa", "Semestre", "Clave", "Nombre", "Cantidad_Grupos",
                "duracion_sesion", "sesiones_semana", "periodo", "udca", "udceq",
                "porcentaje_roster")
    for contexto, r in filas(libro, "Carreras", columnas):
        programa, clave = str(r["Programa"]), str(r["Clave"])
        semestre = numero(r["Semestre"], contexto, entero=True, minimo=1)
        parciales = periodos(r["periodo"], contexto)
        identidad = (programa, semestre, clave, tuple(parciales))
        if identidad in solicitudes:
            raise ValueError(f"{contexto}: solicitud duplicada {identidad}")
        solicitudes.add(identidad)
        grupos = numero(r["Cantidad_Grupos"], contexto, entero=True)
        if grupos == 0:
            continue
        materia = {
            "id": clave,
            "major": programa,
            "semester": semestre,
            "name": str(r["Nombre"]),
            "partials": parciales,
            "load": numero(r["udceq"], contexto),
            "udca": numero(r["udca"], contexto),
            "porcentaje_roster": numero(r["porcentaje_roster"], contexto),
            "sessions_per_week": numero(r["sesiones_semana"], contexto, True, 1),
            "duration_minutes": numero(r["duracion_sesion"], contexto, True, 1),
        }
        materias.append(materia)
        carreras.append({"major": programa, "semester": semestre, "class_id": clave,
                         "cantidad_grupos": grupos, "partials": parciales})

    profesores, disponibilidad = {}, {}
    for contexto, r in filas(libro, "Profesores - Info", ("nomina", "nombre", "limite_uf")):
        nomina = str(r["nomina"])
        if nomina in profesores:
            raise ValueError(f"{contexto}: nómina duplicada {nomina}")
        profesores[nomina] = {
            "id": nomina, "name": str(r["nombre"]), "can_teach": [],
            "max_load_total": numero(r["limite_uf"], contexto), "availability": [],
        }
        disponibilidad[nomina] = {d: [(420, 1260)] for d in dias_semana}

    for contexto, r in filas(libro, "Profesores - Materias", ("nomina", "clave_materia")):
        nomina, clave = str(r["nomina"]), str(r["clave_materia"])
        if nomina not in profesores:
            raise ValueError(f"{contexto}: nómina desconocida {nomina}")
        if clave not in profesores[nomina]["can_teach"]:
            profesores[nomina]["can_teach"].append(clave)

    for contexto, r in filas(libro, "Profesores - Disponibilidad",
                             ("nomina", "Dia", "HoraInicio", "HoraFin")):
        nomina = str(r["nomina"])
        if nomina not in profesores:
            raise ValueError(f"{contexto}: nómina desconocida {nomina}")
        try:
            inicio, fin = minutos(r["HoraInicio"]), minutos(r["HoraFin"])
            dias_bloqueados = dias(r["Dia"])
        except ValueError as error:
            raise ValueError(f"{contexto}: {error}") from None
        if inicio >= fin:
            raise ValueError(f"{contexto}: HoraInicio debe ser menor que HoraFin")
        for d in dias_bloqueados:
            if d in disponibilidad[nomina]:
                disponibilidad[nomina][d] = restar(disponibilidad[nomina][d], inicio, fin)

    for nomina, profesor in profesores.items():
        profesor["can_teach"].sort()
        profesor["availability"] = [
            [d, hora(a), hora(b)]
            for d, intervalos in disponibilidad[nomina].items() for a, b in intervalos
        ]

    aulas = []
    salones = set()
    for contexto, r in filas(libro, "Aulas", ("Salón", "Edificio")):
        salon = str(r["Salón"])
        if salon.casefold() in salones:
            raise ValueError(f"{contexto}: salón duplicado {salon}")
        salones.add(salon.casefold())
        aulas.append({"id": salon, "building": str(r["Edificio"])})
    return {"materias": materias, "carreras": carreras,
            "profesores": list(profesores.values()), "aulas": aulas}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel", type=Path)
    parser.add_argument("--salida", type=Path, default=Path("data"))
    parser.add_argument("--dias", default="LuMaMiJuViSaDo",
                        help="Días de la semana disponibles por defecto")
    args = parser.parse_args()
    try:
        if args.excel.suffix.lower() != ".xlsx":
            raise ValueError("El archivo debe ser .xlsx")
        libro = load_workbook(args.excel, data_only=True)
        try:
            datos = extraer(libro, dias(args.dias))
        finally:
            libro.close()
        # Primero se valida todo. Los errores de datos no producen JSON parciales.
        args.salida.mkdir(parents=True, exist_ok=True)
        for nombre, registros in datos.items():
            destino = args.salida / f"{nombre}.json"
            destino.write_text(json.dumps(registros, ensure_ascii=False,
                                          indent=2, allow_nan=False) + "\n", encoding="utf-8")
            print(f"{destino}: {len(registros)} registros")
    except (ValueError, OSError) as error:
        parser.exit(1, f"Error: {error}\n")


if __name__ == "__main__":
    main()
